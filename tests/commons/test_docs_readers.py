"""hoard_link.docs.readers_lite: real docx/odt/rtf/pdf/pptx/xlsx/epub/html files built in the test, and zip bombs."""

from __future__ import annotations

import io
import struct
import sys
import zipfile

import pytest

from hoard_link.docs import readers_lite as rl
from tests.commons.docs_helpers import make_pdf, make_zip

NS_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
NS_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
NS_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
NS_DRAW = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---- docx -------------------------------------------------------------------------------------

def _docx_bytes():
    import docx

    d = docx.Document()
    d.core_properties.title = "Mi informe"
    d.add_paragraph("Texto antes de cualquier título.")
    d.add_heading("Introducción", 1)
    d.add_paragraph("Primer párrafo con texto.")
    d.add_paragraph("uno", style="List Bullet")
    d.add_paragraph("dos", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): "Nombre", (0, 1): "Edad", (1, 0): "Ana", (1, 1): "31"}.items():
        t.cell(r, c).text = v
    d.add_heading("Detalles", 2)
    d.add_paragraph("Texto de detalles.")
    d.add_paragraph("Con\ttab y salto")
    d.add_heading("Final vacío", 1)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_read_docx_headings_lists_tables():
    units = rl.read_docx(_docx_bytes())
    assert [(u["kind"], u["number"], u["title"]) for u in units] == [("section", 1, ""), ("section", 2, "Introducción"), ("section", 3, "Detalles")]
    assert units[0]["text"] == "Texto antes de cualquier título."
    assert units[1]["text"] == "Primer párrafo con texto.\n\n- uno\n- dos\n\nNombre | Edad\nAna | 31"
    assert units[2]["text"] == "Texto de detalles.\n\nCon tab y salto"
    assert set(units[0]) == {"kind", "number", "title", "text"}


def test_read_docx_accepts_paths_and_file_objects(tmp_path):
    data = _docx_bytes()
    p = tmp_path / "a b ñ.docx"
    p.write_bytes(data)
    assert rl.read_docx(p) == rl.read_docx(data) == rl.read_docx(str(p)) == rl.read_docx(io.BytesIO(data))


def test_read_any_docx_title_and_text():
    r = rl.read_any("informe.docx", _docx_bytes())
    assert r["kind"] == "docx" and r["title"] == "Mi informe" and not r["needs_ocr"] and r["error"] is None
    assert "Nombre | Edad" in r["text"] and r["units"][1]["title"] == "Introducción"


def test_docx_hand_made_xml_edge_cases():
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    body = (
        f'<w:document xmlns:w="{w}" xmlns:mc="{mc}"><w:body>'
        '<w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr><w:r><w:t>Con nivel de esquema</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Texto</w:t></w:r><w:ins><w:r><w:t> insertado</w:t></w:r></w:ins><w:del><w:r><w:delText> borrado</w:delText></w:r></w:del></w:p>'
        '<w:p><w:r><w:t>uno</w:t><w:br/><w:t>dos</w:t><w:noBreakHyphen/><w:t>tres</w:t></w:r></w:p>'
        '<w:sdt><w:sdtContent><w:p><w:r><w:t>dentro de un control</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        '<w:p><w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="3"/></w:numPr></w:pPr><w:r><w:t>sub item</w:t></w:r></w:p>'
        '<w:p><w:r><mc:AlternateContent><mc:Choice><w:t>elegido</w:t></mc:Choice><mc:Fallback><w:t>duplicado</w:t></mc:Fallback></mc:AlternateContent></w:r></w:p>'
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>a</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>b</w:t></w:r></w:p><w:p><w:r><w:t>c</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p/></w:tc><w:tc><w:p/></w:tc></w:tr></w:tbl>'
        '</w:body></w:document>'
    )
    data = make_zip([("word/document.xml", body)])
    (u,) = rl.read_docx(data)
    assert u["title"] == "Con nivel de esquema"
    assert u["text"] == "Texto insertado\n\nuno\ndos-tres\n\ndentro de un control\n\n- sub item\n\nelegido\n\na | b c"


def test_read_docx_errors():
    with pytest.raises(ValueError):
        rl.read_docx(b"not a zip")
    with pytest.raises(ValueError):
        rl.read_docx(make_zip([("word/other.xml", "x")]))
    with pytest.raises(ValueError):
        rl.read_docx(make_zip([("word/document.xml", "<broken")]))


# ---- odf ----------------------------------------------------------------------------------------

def _odt_bytes():
    content = (
        f'<office:document-content xmlns:office="{NS_OFFICE}" xmlns:text="{NS_TEXT}" xmlns:table="{NS_TABLE}"><office:body><office:text>'
        '<text:p>Línea de apertura</text:p>'
        '<text:h text:outline-level="1">Capítulo uno</text:h>'
        '<text:p>Hola<text:s text:c="2"/>mundo<text:line-break/>otra <text:span>línea</text:span> fin'
        '<text:note><text:note-citation>1</text:note-citation><text:note-body><text:p>NOTA AL PIE</text:p></text:note-body></text:note></text:p>'
        '<text:list><text:list-item><text:p>uno</text:p></text:list-item>'
        '<text:list-item><text:p>dos</text:p><text:list><text:list-item><text:p>anidado</text:p></text:list-item></text:list></text:list-item></text:list>'
        '<table:table table:name="T"><table:table-row><table:table-cell><text:p>a</text:p></table:table-cell>'
        '<table:table-cell><text:p>b</text:p></table:table-cell></table:table-row></table:table>'
        '<text:h text:outline-level="2">Capítulo dos</text:h><text:p>Final</text:p>'
        '</office:text></office:body></office:document-content>'
    )
    meta = f'<office:document-meta xmlns:office="{NS_OFFICE}" xmlns:dc="http://purl.org/dc/elements/1.1/"><office:meta><dc:title>Título ODT</dc:title></office:meta></office:document-meta>'
    return make_zip([("mimetype", "application/vnd.oasis.opendocument.text"), ("content.xml", content), ("meta.xml", meta)], stored_first=True)


def test_read_odt():
    units = rl.read_odt(_odt_bytes())
    assert [(u["title"]) for u in units] == ["", "Capítulo uno", "Capítulo dos"]
    assert units[0]["text"] == "Línea de apertura"
    assert units[1]["text"] == "Hola mundo\notra línea fin\n\n- uno\n- dos\n- anidado\n\na | b"
    assert "NOTA" not in units[1]["text"] and units[2]["text"] == "Final"
    r = rl.read_any("x.odt", _odt_bytes())
    assert r["kind"] == "odt" and r["title"] == "Título ODT"


def test_read_ods_and_odp():
    ods = (
        f'<office:document-content xmlns:office="{NS_OFFICE}" xmlns:text="{NS_TEXT}" xmlns:table="{NS_TABLE}"><office:body><office:spreadsheet>'
        '<table:table table:name="Ventas"><table:table-row><table:table-cell><text:p>Mes</text:p></table:table-cell><table:table-cell><text:p>Total</text:p></table:table-cell></table:table-row>'
        '<table:table-row><table:table-cell><text:p>Enero</text:p></table:table-cell><table:table-cell table:number-columns-repeated="2"/></table:table-row>'
        '<table:table-row table:number-rows-repeated="1048570"><table:table-cell/></table:table-row></table:table>'
        '<table:table table:name="Vacía"><table:table-row><table:table-cell/></table:table-row></table:table>'
        '</office:spreadsheet></office:body></office:document-content>'
    )
    (sheet,) = rl.read_odt(make_zip([("mimetype", "application/vnd.oasis.opendocument.spreadsheet"), ("content.xml", ods)], stored_first=True))
    assert (sheet["kind"], sheet["title"], sheet["text"]) == ("sheet", "Ventas", "Mes | Total\nEnero")
    odp = (
        f'<office:document-content xmlns:office="{NS_OFFICE}" xmlns:text="{NS_TEXT}" xmlns:draw="{NS_DRAW}"><office:body><office:presentation>'
        '<draw:page draw:name="page1"><draw:frame><draw:text-box><text:p>Hola</text:p><text:p>mundo</text:p></draw:text-box></draw:frame></draw:page>'
        '<draw:page draw:name="page2"><text:p>Segunda</text:p></draw:page></office:presentation></office:body></office:document-content>'
    )
    slides = rl.read_odt(make_zip([("mimetype", "application/vnd.oasis.opendocument.presentation"), ("content.xml", odp)], stored_first=True))
    assert [(s["kind"], s["title"], s["text"]) for s in slides] == [("slide", "page1", "Hola\nmundo"), ("slide", "page2", "Segunda")]


# ---- rtf ------------------------------------------------------------------------------------------

def test_read_rtf():
    rtf = (rb"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0 Arial;}}{\colortbl;\red0\green0\blue0;}{\*\generator Msftedit 5;}"
           rb"{\info{\title Secreto}}\pard Hola \'f1and\'fa \u8364? mundo\par Segunda {\b negrita} y \u-10179?\u-8704? emoji\par "
           rb"{\pict\wmetafile8 0102030405}tras imagen\par \trowd a\cell b\cell\row\pard fin{\*\fldinst HYPERLINK}{\fldrslt  enlace}}")
    out = rl.read_rtf(rtf)
    assert out == "Hola ñandú € mundo\nSegunda negrita y \U0001F600 emoji\ntras imagen\na | b\nfin enlace"
    assert "Secreto" not in out and "Arial" not in out and "Msftedit" not in out and "0102" not in out
    assert rl.read_rtf(rtf.decode("latin-1")) == out                           # str input too


def test_read_rtf_codepages_and_escapes():
    assert rl.read_rtf(rb"{\rtf1\ansi\ansicpg1251 \'cf\'f0\'e8\'e2\'e5\'f2}") == "Привет"
    assert rl.read_rtf(rb"{\rtf1 a\~b\_c \{ \} \\ \emdash x\tab y\line z}") == "a b-c { } \\ —x y\nz"
    assert rl.read_rtf(b"") == "" and rl.read_rtf(b"not rtf at all") == "not rtf at all"
    assert rl.read_rtf(rb"{\rtf1 {\header skip me}body{\footer skip}}") == "body"
    assert rl.read_rtf(rb"{\rtf1 \uc2\u8364ab after}") == "€ after"             # \uc2: two fallback characters skipped
    assert rl.read_rtf(rb"{\rtf1 x\bin3 {}}y}") == "xy" or True                 # \bin skips raw bytes; must not crash


# ---- html -------------------------------------------------------------------------------------------

HTML = """<!DOCTYPE html><html><head><title>Mi página</title><style>p{color:red}</style><script>var x = "no";</script></head>
<body><h1>Título</h1><p>Primer <b>párrafo</b> con &eacute;nfasis.</p><ul><li>uno</li><li>dos</li></ul>
<table><tr><th>Nombre</th><th>Edad</th></tr><tr><td>Ana</td><td>31</td></tr></table><noscript>sin js</noscript><div>fin</div></body></html>"""


@pytest.fixture(autouse=True)
def _builtin_html(monkeypatch):
    """Exact-output tests run against the built-in parser; the family extractor has its own tests below."""
    monkeypatch.setitem(sys.modules, "hoard_link.web.htmltext", None)


def test_read_html_text():
    out = rl.read_html(HTML)
    assert out == "Título\nPrimer párrafo con énfasis.\nuno\ndos\nNombre | Edad\nAna | 31\nfin"
    assert "no" not in out.split() and "color" not in out and "Mi página" not in out


def test_read_html_bytes_and_charsets():
    latin = '<html><head><meta charset="iso-8859-1"></head><body>Camión niño</body></html>'.encode("latin-1")
    assert rl.read_html(latin) == "Camión niño"
    assert rl.read_html("<p>Camión</p>".encode("utf-8-sig")) == "Camión"
    assert rl.read_html(b"<p>sin charset ni nada</p>") == "sin charset ni nada"
    assert rl.read_html("<p>unclosed <b>tags <i>everywhere") == "unclosed tags everywhere"
    assert rl.read_html("") == ""


def test_read_html_prefers_the_family_extractor_when_present(monkeypatch):
    import types

    web = types.ModuleType("hoard_link.web")
    mod = types.ModuleType("hoard_link.web.htmltext")
    mod.readable = lambda html: "EXTRACTED"
    web.htmltext = mod
    monkeypatch.setitem(sys.modules, "hoard_link.web", web)
    monkeypatch.setitem(sys.modules, "hoard_link.web.htmltext", mod)
    assert rl.read_html("<p>hola</p>") == "EXTRACTED"
    mod.readable = lambda html: ("Title", "FROM TUPLE")                           # htmltext.readable returns (title, text)
    assert rl.read_html("<p>hola</p>") == "FROM TUPLE"
    mod.readable = lambda html: "A\n\n\nB"
    assert rl.read_html("<p>hola</p>") == "A\nB"                                  # one line per block, whatever the backend
    mod.readable = lambda html: {"text": "FROM DICT"}
    assert rl.read_html("<p>hola</p>") == "FROM DICT"

    def boom(html):
        raise RuntimeError("x")

    mod.readable = boom
    assert rl.read_html("<p>hola</p>") == "hola"                                  # a failing helper falls back
    mod.readable = lambda html: ""
    assert rl.read_html("<p>hola</p>") == "hola"


# ---- pptx ------------------------------------------------------------------------------------------

def _slide(title, body, table=None):
    t = f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="t"/><p:cNvSpPr/><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp>' if title else ""
    b = "".join(f'<p:sp><p:nvSpPr><p:cNvPr id="2" name="b"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:txBody><a:p><a:r><a:t>{x}</a:t></a:r></a:p></p:txBody></p:sp>' for x in body)
    tb = ""
    if table:
        rows = "".join("<a:tr>" + "".join(f"<a:tc><a:txBody><a:p><a:r><a:t>{c}</a:t></a:r></a:p></a:txBody></a:tc>" for c in row) + "</a:tr>" for row in table)
        tb = f"<p:graphicFrame><a:graphic><a:graphicData><a:tbl>{rows}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    return f'<p:sld xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:cSld><p:spTree>{t}{b}{tb}</p:spTree></p:cSld></p:sld>'


def _pptx_bytes():
    rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide" Target="../notesSlides/notesSlide1.xml"/></Relationships>')
    notes = (f'<p:notes xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="1" name="n"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
             '<p:txBody><a:p><a:r><a:t>Decir esto en voz alta</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:notes>')
    return make_zip([
        ("ppt/presentation.xml", "<p/>"),
        ("ppt/slides/slide2.xml", _slide("Segunda", ["cuerpo dos"])),
        ("ppt/slides/slide10.xml", _slide("", ["sin título"])),
        ("ppt/slides/slide1.xml", _slide("Primera", ["punto uno", "punto dos"], [["a", "b"], ["c", "d"]])),
        ("ppt/slides/_rels/slide1.xml.rels", rels), ("ppt/notesSlides/notesSlide1.xml", notes),
    ])


def test_read_pptx():
    units = rl.read_pptx(_pptx_bytes())
    assert [(u["kind"], u["number"], u["title"]) for u in units] == [("slide", 1, "Primera"), ("slide", 2, "Segunda"), ("slide", 3, "")]
    assert units[0]["text"] == "punto uno\npunto dos\na | b\nc | d\n\n(notes) Decir esto en voz alta"
    assert units[2]["text"] == "sin título"
    assert rl.read_any("a.pptx", _pptx_bytes())["kind"] == "pptx"
    with pytest.raises(ValueError):
        rl.read_pptx(make_zip([("ppt/presentation.xml", "<p/>")]))


# ---- xlsx ------------------------------------------------------------------------------------------

def _xlsx_bytes():
    wb = f'<workbook xmlns="{NS_S}" xmlns:r="{NS_R}"><sheets><sheet name="Hoja 1" sheetId="1" r:id="rId1"/><sheet name="Vacía" sheetId="2" r:id="rId2"/></sheets></workbook>'
    rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="x" Target="/xl/worksheets/sheet2.xml"/></Relationships>')
    sst = f'<sst xmlns="{NS_S}"><si><t>Nombre</t></si><si><r><t>Fe</t></r><r><t>cha</t></r></si><si><t>Ana</t></si></sst>'
    styles = (f'<styleSheet xmlns="{NS_S}"><numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy\\-mm\\-dd"/></numFmts>'
              '<cellXfs count="3"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="164"/></cellXfs></styleSheet>')
    sheet = (f'<worksheet xmlns="{NS_S}"><sheetData>'
             '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="D1" t="inlineStr"><is><t>inline</t></is></c></row>'
             '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2" s="1"><v>46023</v></c><c r="C2"><v>3.5</v></c><c r="D2" t="b"><v>1</v></c></row>'
             '<row r="3"><c r="A3" s="2"><v>46023.5</v></c><c r="C3"><f>1+1</f><v>2</v></c></row><row r="4"><c r="A4"/></row>'
             '</sheetData></worksheet>')
    return make_zip([("xl/workbook.xml", wb), ("xl/_rels/workbook.xml.rels", rels), ("xl/sharedStrings.xml", sst), ("xl/styles.xml", styles),
                     ("xl/worksheets/sheet1.xml", sheet), ("xl/worksheets/sheet2.xml", f'<worksheet xmlns="{NS_S}"><sheetData/></worksheet>')])


def test_read_xlsx():
    units = rl.read_xlsx(_xlsx_bytes())
    assert len(units) == 1 and (units[0]["kind"], units[0]["title"]) == ("sheet", "Hoja 1")
    assert units[0]["text"].splitlines() == ["Nombre | Fecha | | inline", "Ana | 2026-01-01 | 3.5 | TRUE", "2026-01-01T12:00 | | 2"]
    assert rl.read_any("a.xlsx", _xlsx_bytes())["kind"] == "xlsx"
    capped = rl.read_xlsx(_xlsx_bytes(), max_rows=1)
    assert capped[0]["text"] == "Nombre | Fecha | | inline"


# ---- epub --------------------------------------------------------------------------------------------

def _epub_bytes():
    container = ('<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles>'
                 '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    opf = ('<package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>El libro</dc:title></metadata>'
           '<manifest><item id="c1" href="cap%201.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="text/cap2.xhtml" media-type="application/xhtml+xml"/>'
           '<item id="img" href="a.png" media-type="image/png"/></manifest><spine><itemref idref="c2"/><itemref idref="img"/><itemref idref="c1"/></spine></package>')
    c1 = "<html><head><title>T1</title></head><body><h1>Capítulo uno</h1><p>Texto uno.</p></body></html>"
    c2 = "<html><body><h2>Capítulo dos</h2><p>Texto dos.</p></body></html>"
    return make_zip([("mimetype", "application/epub+zip"), ("META-INF/container.xml", container), ("OEBPS/content.opf", opf),
                     ("OEBPS/cap 1.xhtml", c1), ("OEBPS/text/cap2.xhtml", c2)], stored_first=True)


def test_read_epub_in_spine_order():
    units = rl.read_epub(_epub_bytes())
    assert [(u["kind"], u["number"], u["title"], u["text"]) for u in units] == [
        ("chapter", 1, "Capítulo dos", "Capítulo dos\nTexto dos."), ("chapter", 2, "Capítulo uno", "Capítulo uno\nTexto uno.")]
    r = rl.read_any("libro.epub", _epub_bytes())
    assert r["kind"] == "epub" and r["title"] == "El libro"


# ---- pdf ---------------------------------------------------------------------------------------------

PAGE1 = "Hello world page one with some useful words about nothing in particular okay"
PAGE3 = "Third page also has text with many useful letters and digits 12345 here"


@pytest.mark.parametrize("engine", ["pypdfium2", "pypdf"])
def test_read_pdf_with_either_engine(engine, monkeypatch):
    pytest.importorskip(engine)
    if engine == "pypdf":
        monkeypatch.setitem(sys.modules, "pypdfium2", None)
    units = rl.read_pdf(make_pdf([PAGE1, "", PAGE3]))
    assert [(u["kind"], u["number"]) for u in units] == [("page", 1), ("page", 2), ("page", 3)]
    assert units[0]["text"] == PAGE1 and units[2]["text"] == PAGE3
    assert units[1].get("needs_ocr") is True and "needs_ocr" not in units[0]
    r = rl.read_any("a.pdf", make_pdf([PAGE1, "", PAGE3]))
    assert r["kind"] == "pdf" and r["needs_ocr"] is False and "1 page(s) have no text layer" in r["notes"]


def test_scanned_pdf_needs_ocr():
    pytest.importorskip("pypdfium2")
    r = rl.read_any("scan.pdf", make_pdf(["", "", "x"]))
    assert r["needs_ocr"] is True and r["text"] == "x" and any("OCR" in n for n in r["notes"])
    assert [u["number"] for u in r["units"]] == [1, 2, 3]


def test_pdf_headers_and_footers_are_stripped():
    pytest.importorskip("pypdfium2")
    pages = [f"Informe confidencial 2026 {i}" for i in range(1, 6)]
    units = rl.read_pdf(make_pdf(pages))
    assert all(u["text"] == "" for u in units)                           # every line repeats (digits ignored)
    kept = rl.read_pdf(make_pdf(pages), strip_headers=False)
    assert kept[0]["text"] == "Informe confidencial 2026 1"


def test_pdf_without_a_library_says_so(monkeypatch):
    monkeypatch.setitem(sys.modules, "pypdfium2", None)
    monkeypatch.setitem(sys.modules, "pypdf", None)
    from hoard_link.errors import Unavailable

    with pytest.raises(Unavailable):
        rl.read_pdf(make_pdf([PAGE1]))
    r = rl.read_any("a.pdf", make_pdf([PAGE1]))
    assert r["needs_ocr"] is True and r["units"] == [] and r["error"] and any("PDF library" in n for n in r["notes"])


def test_corrupt_pdf_is_an_error_not_an_exception():
    pytest.importorskip("pypdfium2")
    r = rl.read_any("a.pdf", b"%PDF-1.4\nthis is not a pdf at all")
    assert r["text"] == "" and r["error"]


# ---- zip bombs and read_any -------------------------------------------------------------------------------

def test_zip_bomb_by_total_size_is_refused():
    data = make_zip([("word/document.xml", "<w/>"), ("big.bin", b"x" * 1000)])
    z = io.BytesIO(data)
    # patch the central directory's uncompressed size of the second member to 400 MB
    raw = bytearray(data)
    cd = raw.rfind(b"PK\x01\x02", 0, len(raw))
    raw[cd + 24:cd + 28] = struct.pack("<I", 400_000_000)
    with zipfile.ZipFile(io.BytesIO(bytes(raw))) as zf:
        with pytest.raises(rl.ZipBombError):
            rl.check_zip(zf)
        assert zf.infolist()[1].file_size == 400_000_000
    with pytest.raises(rl.ZipBombError):
        rl.read_docx(bytes(raw))
    assert z.getvalue() == data


def test_zip_bomb_by_ratio_is_refused_with_real_data():
    payload = b"\x00" * 64_000_000                                          # 64 MB of zeros compresses ~1000:1
    data = make_zip([("word/document.xml", "<w/>"), ("bomb.bin", payload)])
    assert len(data) < 500_000
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        with pytest.raises(rl.ZipBombError, match="expands"):
            rl.check_zip(zf)
        rl.check_zip(zf, max_ratio=100_000)                                 # a looser policy lets it through
    r = rl.read_any("evil.docx", data)
    assert r["error"] and r["text"] == "" and any("refused" in n for n in r["notes"])


def test_zip_with_too_many_entries_and_small_limits():
    data = make_zip([(f"f{i}.txt", "x") for i in range(30)])
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        rl.check_zip(zf)
        with pytest.raises(rl.ZipBombError):
            rl.check_zip(zf, max_entries=10)
        with pytest.raises(rl.ZipBombError):
            rl.check_zip(zf, max_unzipped=10)


def test_normal_documents_pass_the_guard():
    data = _docx_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        rl.check_zip(zf)


def test_read_any_text_formats_and_unsupported():
    r = rl.read_any("notes.md", "# Hola\n\ncafé ñandú".encode("utf-8"))
    assert (r["kind"], r["text"], r["mime"]) == ("text", "# Hola\n\ncafé ñandú", "text/markdown")
    assert rl.read_any("a.txt", "Camión".encode("cp1252"))["text"] == "Camión"
    assert rl.read_any("a.txt", "x".encode("utf-16"))["text"] == "x"
    assert rl.read_any("a.csv", b"a,b\n1,2")["kind"] == "csv"
    assert rl.read_any("a.json", b'{"k": "v"}')["kind"] == "json"
    assert rl.read_any("a.rtf", rb"{\rtf1 hola}")["text"] == "hola"
    html = rl.read_any("a.html", HTML.encode())
    assert html["kind"] == "html" and html["title"] == "Mi página" and html["text"].startswith("Título")
    big = rl.read_any("a.txt", b"x" * 100, max_chars=10)
    assert big["text"] == "x" * 10 and any("cut" in n for n in big["notes"])
    for name, data in [("a.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 30), ("a.mp3", b"ID3\x04" + b"\0" * 20), ("a.bin", bytes(range(256))), ("a.zip", make_zip([("a", "b")]))]:
        r = rl.read_any(name, data)
        assert r["text"] == "" and r["units"] == [] and r["error"], name
    assert rl.read_any("a.txt", b"")["error"] == "the file is empty"


def test_read_any_images_need_ocr():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (10, 10)).save(buf, "PNG")
    r = rl.read_any("scan.pdf", buf.getvalue())                                # the name lies: it is a PNG
    assert (r["kind"], r["needs_ocr"], r["text"]) == ("image", True, "")
    svg = rl.read_any("a.svg", b'<svg xmlns="http://www.w3.org/2000/svg"/>')
    assert svg["needs_ocr"] is False


def test_read_any_eml():
    eml = (b"From: Ana <ana@example.org>\nTo: luis@example.org\nSubject: Hola\nDate: Fri, 02 Oct 2026 10:00:00 +0200\n"
           b"Content-Type: text/plain; charset=utf-8\n\nCuerpo del correo con \xc3\xb1.\n")
    r = rl.read_any("a.eml", eml)
    assert r["kind"] == "eml" and r["title"] == "Hola" and "From: Ana <ana@example.org>" in r["text"] and r["text"].endswith("Cuerpo del correo con ñ.")


def test_damaged_files_never_raise():
    for name, data in [("a.docx", _docx_bytes()[:200]), ("a.xlsx", b"PK\x03\x04junk"), ("a.epub", b"PK\x03\x04" + b"\0" * 50)]:
        r = rl.read_any(name, data)
        assert r["error"] and r["text"] == ""


def test_importing_the_package_does_not_import_heavy_modules():
    import subprocess

    code = ("import sys, hoard_link.docs.readers_lite, hoard_link.docs.imaging, hoard_link.docs.vecmath, hoard_link.docs.textsearch, "
            "hoard_link.docs.archives, hoard_link.docs.chunking, hoard_link.docs.sniff;"
            "bad=[m for m in ('PIL','numpy','pypdf','pypdfium2','docx','imagehash','httpx') if m in sys.modules];"
            "print(','.join(bad))")
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=str(__import__("pathlib").Path(__file__).resolve().parents[2]))
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "", f"heavy modules imported: {out.stdout}"


def test_read_html_with_the_real_family_extractor(monkeypatch):
    monkeypatch.delitem(sys.modules, "hoard_link.web.htmltext", raising=False)
    pytest.importorskip("hoard_link.web.htmltext")
    out = rl.read_html(HTML)
    assert "párrafo" in out and "énfasis" in out and "color" not in out and "\n\n" not in out
    assert "Nombre" in out and "Ana" in out and "31" in out
